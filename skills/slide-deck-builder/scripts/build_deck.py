#!/usr/bin/env python3
"""Build a 16:9 .pptx deck from an approved outline (JSON) with python-pptx.

Usage:
    python3 scripts/build_deck.py outline.json deck.pptx [--brand brand.json]

The outline format is documented in references/outline-format.md. The brand file is optional;
without it the deck uses a neutral palette. After building, the script prints a list of
warnings (likely text overflow, too many bullets, long headlines). Fix every warning in the
outline and rebuild before delivery.

Requires: pip install python-pptx
"""
import argparse
import json
import math
import os
import sys

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt
except ImportError:
    sys.exit("python-pptx is missing: pip install python-pptx")

NEUTRAL = {
    "background": "#FFFFFF",
    "ink": "#1A1A1A",
    "body": "#3A3A3A",
    "muted": "#6B6B6B",
    "primary": "#1F3A5F",
    "accent": "#2F7D6D",
    "surface": "#F4F4F2",
    "font_heading": "Arial",
    "font_body": "Arial",
    "logo": None,
    "footer": "",
}

W, H = 13.333, 7.5          # slide size, inches
MARGIN = 0.6
WARNINGS = []


def rgb(hex_str):
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def estimate_overflow(slide_no, label, text, font_pt, box_w, box_h):
    """Rough fit check: average glyph width ~0.5 em, line height ~1.25 em."""
    if not text:
        return
    char_w = 0.5 * font_pt / 72.0
    per_line = max(1, int(box_w / char_w))
    lines = sum(max(1, math.ceil(len(p) / per_line)) for p in text.split("\n"))
    need = lines * 1.25 * font_pt / 72.0
    if need > box_h:
        WARNINGS.append(
            f"slide {slide_no}: {label} likely overflows ({lines} lines at {font_pt}pt need "
            f"{need:.1f}in, box is {box_h:.1f}in). Shorten the text or split the slide."
        )


def add_text(slide, slide_no, label, text, x, y, w, h, size, color, font, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, para in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = para
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = rgb(color)
    estimate_overflow(slide_no, label, str(text), size, w - 0.1, h)
    return box


def add_bullets(slide, slide_no, items, x, y, w, h, b, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = "•  " + item
        run.font.size = Pt(size)
        run.font.name = b["font_body"]
        run.font.color.rgb = rgb(b["body"])
    estimate_overflow(slide_no, "bullets", "\n".join("xx" + i for i in items), size, w - 0.1, h)
    if len(items) > 5:
        WARNINGS.append(f"slide {slide_no}: {len(items)} bullets; keep to 5 or fewer, or split.")


def fill_bg(slide, b):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(b["background"])


def accent_bar(slide, b, x=MARGIN, y=0.55, w=0.9, h=0.08):
    bar = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(b["accent"])
    bar.line.fill.background()


def footer(slide, b, n, total):
    if b.get("footer"):
        add_text(slide, n, "footer", b["footer"], MARGIN, H - 0.5, 8, 0.3, 10, b["muted"],
                 b["font_body"])
    add_text(slide, n, "page", f"{n}/{total}", W - MARGIN - 1, H - 0.5, 1, 0.3, 10,
             b["muted"], b["font_body"], align=PP_ALIGN.RIGHT)
    logo = b.get("logo")
    if logo and os.path.exists(logo):
        slide.shapes.add_picture(logo, Inches(W - MARGIN - 1.2), Inches(0.35), height=Inches(0.45))


def headline(slide, n, s, b, y=0.75, h=1.3, size=30):
    text = s.get("headline", "")
    words = len(text.split())
    if words > 16:
        WARNINGS.append(f"slide {n}: headline has {words} words; aim for 8-14.")
    add_text(slide, n, "headline", text, MARGIN, y, W - 2 * MARGIN, h, size, b["ink"],
             b["font_heading"], bold=True)


def build(outline, brand, out_path):
    b = dict(NEUTRAL)
    b.update({k: v for k, v in (brand or {}).items() if v is not None})
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    slides = outline.get("slides", [])
    total = len(slides)

    for n, s in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        fill_bg(slide, b)
        kind = s.get("layout", "bullets")
        body_top = 2.2

        if kind in ("title", "closing"):
            band = slide.shapes.add_shape(1, 0, 0, Inches(W), Inches(H))
            band.fill.solid()
            band.fill.fore_color.rgb = rgb(b["primary"])
            band.line.fill.background()
            add_text(slide, n, "headline", s.get("headline", ""), MARGIN + 0.2, 2.2,
                     W - 2 * MARGIN - 0.4, 2.0, 40, "#FFFFFF", b["font_heading"], bold=True)
            if s.get("sub"):
                add_text(slide, n, "sub", s["sub"], MARGIN + 0.2, 4.4, W - 2 * MARGIN - 0.4,
                         1.2, 20, "#FFFFFF", b["font_body"])
        elif kind == "section":
            accent_bar(slide, b, y=3.0)
            add_text(slide, n, "headline", s.get("headline", ""), MARGIN, 3.2,
                     W - 2 * MARGIN, 1.5, 36, b["ink"], b["font_heading"], bold=True)
        else:
            accent_bar(slide, b)
            headline(slide, n, s, b)
            if s.get("sub"):
                add_text(slide, n, "sub", s["sub"], MARGIN, 1.9, W - 2 * MARGIN, 0.5, 16,
                         b["muted"], b["font_body"])
                body_top = 2.5
            body_h = H - body_top - 0.8

            if kind == "bullets":
                add_bullets(slide, n, s.get("bullets", []), MARGIN, body_top,
                            W - 2 * MARGIN, body_h, b)
            elif kind == "stat":
                stats = s.get("stats", [])[:4]
                if stats:
                    col_w = (W - 2 * MARGIN - 0.3 * (len(stats) - 1)) / len(stats)
                    for i, st in enumerate(stats):
                        x = MARGIN + i * (col_w + 0.3)
                        card = slide.shapes.add_shape(1, Inches(x), Inches(body_top),
                                                      Inches(col_w), Inches(2.6))
                        card.fill.solid()
                        card.fill.fore_color.rgb = rgb(b["surface"])
                        card.line.fill.background()
                        add_text(slide, n, "stat value", st.get("value", ""), x + 0.2,
                                 body_top + 0.25, col_w - 0.4, 1.1, 44, b["primary"],
                                 b["font_heading"], bold=True)
                        add_text(slide, n, "stat label", st.get("label", ""), x + 0.2,
                                 body_top + 1.4, col_w - 0.4, 1.1, 16, b["body"],
                                 b["font_body"])
                if s.get("source"):
                    add_text(slide, n, "source", "Source: " + s["source"], MARGIN, H - 1.0,
                             W - 2 * MARGIN, 0.35, 11, b["muted"], b["font_body"])
                else:
                    WARNINGS.append(f"slide {n}: stat slide without a source.")
            elif kind == "two_col":
                col_w = (W - 2 * MARGIN - 0.5) / 2
                for i, side in enumerate(("left", "right")):
                    col = s.get(side, {})
                    x = MARGIN + i * (col_w + 0.5)
                    add_text(slide, n, f"{side} title", col.get("title", ""), x, body_top,
                             col_w, 0.5, 20, b["accent"], b["font_heading"], bold=True)
                    add_bullets(slide, n, col.get("bullets", []), x, body_top + 0.6, col_w,
                                body_h - 0.6, b, size=18)
            elif kind == "quote":
                add_text(slide, n, "quote", "“" + s.get("quote", "") + "”", MARGIN + 0.5,
                         body_top, W - 2 * MARGIN - 1, body_h - 0.8, 26, b["ink"],
                         b["font_heading"])
                add_text(slide, n, "attribution", s.get("attribution", ""), MARGIN + 0.5,
                         H - 1.5, W - 2 * MARGIN - 1, 0.4, 14, b["muted"], b["font_body"])
            elif kind == "table":
                rows = s.get("rows", [])
                if rows:
                    nr, nc = len(rows), max(len(r) for r in rows)
                    if nr > 8:
                        WARNINGS.append(f"slide {n}: table has {nr} rows; keep to 8 or move it to an appendix.")
                    tbl = slide.shapes.add_table(nr, nc, Inches(MARGIN), Inches(body_top),
                                                 Inches(W - 2 * MARGIN),
                                                 Inches(min(body_h, 0.45 * nr))).table
                    for r, row in enumerate(rows):
                        for c in range(nc):
                            cell = tbl.cell(r, c)
                            cell.text = str(row[c]) if c < len(row) else ""
                            for p in cell.text_frame.paragraphs:
                                for run in p.runs:
                                    run.font.size = Pt(14)
                                    run.font.name = b["font_body"]
                                    run.font.bold = r == 0
                                    run.font.color.rgb = rgb("#FFFFFF" if r == 0 else b["body"])
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = rgb(b["primary"] if r == 0 else b["background"])
            else:
                WARNINGS.append(f"slide {n}: unknown layout '{kind}', rendered as headline only.")

        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
        if kind not in ("title", "closing"):
            footer(slide, b, n, total)

    prs.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("outline")
    ap.add_argument("out")
    ap.add_argument("--brand", help="brand JSON (colors, fonts, logo, footer)")
    a = ap.parse_args()
    outline = json.load(open(a.outline, encoding="utf-8"))
    brand = json.load(open(a.brand, encoding="utf-8")) if a.brand else outline.get("brand")
    build(outline, brand, a.out)
    print(f"built {a.out} ({len(outline.get('slides', []))} slides)")
    if WARNINGS:
        print(f"{len(WARNINGS)} warning(s):")
        for w in WARNINGS:
            print("  - " + w)
    else:
        print("no warnings")


if __name__ == "__main__":
    main()
